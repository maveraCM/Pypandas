import sqlite3
import pandas as pd
import os
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference

def export_to_excel_with_chart(table_name):
    db_path = '/Users/user/Desktop/pyproj/testdb.sql'
    
    try:
        # Проверяем существование БД
        if not os.path.exists(db_path):
            print(f"❌ База данных не найдена: {db_path}")
            return
            
        # Подключаемся к БД
        conn = sqlite3.connect(db_path)
        
        # Получаем данные
        query = f"SELECT * FROM {table_name}"
        df = pd.read_sql_query(query, conn)
        
        if df.empty:
            print(f"📝 Таблица {table_name} пуста")
            return
            
        # Создаем имя файла с текущей датой
        current_time = datetime.now().strftime("%Y%m%d_%H%M%S")
        excel_path = f'/Users/user/Desktop/pyproj/export_{current_time}.xlsx'
        
        # Экспортируем в Excel
        df.to_excel(excel_path, sheet_name='Data', index=False)
        
        # Создаем диаграмму
        wb = load_workbook(excel_path)
        ws = wb['Data']
        
        # Создаем график (настройте под свои данные)
        chart = BarChart()
        chart.title = f"Данные из таблицы {table_name}"
        
        # Получаем данные для графика (замените столбцы на свои)
        data = Reference(ws, min_col=2, max_col=2, min_row=1, max_row=len(df)+1)
        cats = Reference(ws, min_col=1, max_col=1, min_row=2, max_row=len(df)+1)
        
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        
        # Добавляем график на лист
        ws.add_chart(chart, "H2")
        
        # Сохраняем файл
        wb.save(excel_path)
        
        print(f"✅ Данные экспортированы в файл: {excel_path}")
        print(f"📊 График создан")
            
    except Exception as e:
        print(f"❌ Ошибка: {e}")
    finally:
        if 'conn' in locals():
            conn.close()

if __name__ == "__main__":
    table_name = 'orders'
    export_to_excel_with_chart(table_name)
